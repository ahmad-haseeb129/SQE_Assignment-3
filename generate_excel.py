"""
Generate test_cases.xlsx for SQE Assignment 3
Run: python generate_excel.py
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()

# --- Styles ---
header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
pass_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
fail_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

def style_header(ws, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

def style_rows(ws, rows, cols):
    for row in range(2, rows + 2):
        for col in range(1, cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True, vertical='top')

# ===================== Sheet 1: Parser Tests =====================
ws1 = wb.active
ws1.title = "Parser Tests"
headers1 = ["Test ID", "Module", "Test Type", "Test Description", "Input", "Parameters", "Expected Output", "Status", "Author"]
for i, h in enumerate(headers1, 1):
    ws1.cell(row=1, column=i, value=h)

parser_tests = [
    ["TC_P01", "parser", "ECP", "Parse ISO 8601 format", "2024-01-15", "default", "datetime(2024,1,15)", "Pass", "Ibrahim"],
    ["TC_P02", "parser", "ECP", "Parse US date format", "01/15/2024", "default", "datetime(2024,1,15)", "Pass", "Ibrahim"],
    ["TC_P03", "parser", "ECP", "Parse written month name", "January 15, 2024", "default", "datetime(2024,1,15)", "Pass", "Ibrahim"],
    ["TC_P04", "parser", "ECP", "Parse abbreviated month", "Jan 15, 2024", "default", "datetime(2024,1,15)", "Pass", "Ibrahim"],
    ["TC_P05", "parser", "ECP", "Parse date with time", "2024-01-15 14:30:00", "default", "datetime(2024,1,15,14,30)", "Pass", "Ibrahim"],
    ["TC_P06", "parser", "ECP", "dayfirst=True on ambiguous", "01/02/2024", "dayfirst=True", "datetime(2024,2,1)", "Pass", "Ibrahim"],
    ["TC_P07", "parser", "BVA", "Min year boundary", "0001-01-01", "default", "datetime(1,1,1)", "Pass", "Ibrahim"],
    ["TC_P08", "parser", "BVA", "Max year boundary", "9999-12-31", "default", "datetime(9999,12,31)", "Pass", "Ibrahim"],
    ["TC_P09", "parser", "BVA", "Leap year Feb 29", "2024-02-29", "default", "datetime(2024,2,29)", "Pass", "Ibrahim"],
    ["TC_P10", "parser", "ECP", "Empty string raises error", "", "default", "ParserError", "Pass", "Ahmad"],
    ["TC_P11", "parser", "ECP", "Non-date text raises error", "hello world", "default", "ParserError", "Pass", "Ahmad"],
    ["TC_P12", "parser", "ECP", "None input raises TypeError", "None", "default", "TypeError", "Pass", "Ahmad"],
    ["TC_P13", "parser", "ECP", "Fuzzy parsing with noise", "The date is Jan 15, 2024 ok", "fuzzy=True", "datetime(2024,1,15)", "Pass", "Ahmad"],
    ["TC_P14", "parser", "BVA", "Time boundary midnight", "2024-01-15 00:00:00", "default", "datetime(2024,1,15,0,0,0)", "Pass", "Ahmad"],
    ["TC_P15", "parser", "BVA", "Time boundary end of day", "2024-01-15 23:59:59", "default", "datetime(2024,1,15,23,59,59)", "Pass", "Ahmad"],
]

for i, row in enumerate(parser_tests, 2):
    for j, val in enumerate(row, 1):
        ws1.cell(row=i, column=j, value=val)

style_header(ws1, len(headers1))
style_rows(ws1, len(parser_tests), len(headers1))
for col in ws1.columns:
    ws1.column_dimensions[col[0].column_letter].width = 22

# ===================== Sheet 2: Relativedelta Tests =====================
ws2 = wb.create_sheet("Relativedelta Tests")
headers2 = ["Test ID", "Module", "Test Type", "Test Description", "Input Date", "Delta", "Expected Output", "Status", "Author"]
for i, h in enumerate(headers2, 1):
    ws2.cell(row=1, column=i, value=h)

rd_tests = [
    ["TC_RD01", "relativedelta", "ECP", "Positive year offset", "2024-01-15", "years=+1", "2025-01-15", "Pass", "Ibrahim"],
    ["TC_RD02", "relativedelta", "ECP", "Negative year offset", "2024-06-15", "years=-1", "2023-06-15", "Pass", "Ibrahim"],
    ["TC_RD03", "relativedelta", "ECP", "Positive month offset", "2024-01-15", "months=+3", "2024-04-15", "Pass", "Ibrahim"],
    ["TC_RD04", "relativedelta", "ECP", "Combined offsets", "2024-01-15", "years=1,months=2,days=5", "2025-03-20", "Pass", "Ibrahim"],
    ["TC_RD05", "relativedelta", "BVA", "Month-end overflow (leap)", "2024-01-31", "months=+1", "2024-02-29", "Pass", "Ibrahim"],
    ["TC_RD06", "relativedelta", "BVA", "Month-end overflow (non-leap)", "2023-01-31", "months=+1", "2023-02-28", "Pass", "Ibrahim"],
    ["TC_RD07", "relativedelta", "BVA", "Leap year +1 year", "2024-02-29", "years=+1", "2025-02-28", "Pass", "Ibrahim"],
    ["TC_RD08", "relativedelta", "ECP", "Multiply relativedelta", "rd(months=2,days=3)", "* 3", "rd(months=6,days=9)", "Pass", "Ibrahim"],
    ["TC_RD09", "relativedelta", "ECP", "Absolute month (singular)", "2024-06-15", "month=1", "2024-01-15", "Pass", "Ahmad"],
    ["TC_RD10", "relativedelta", "ECP", "Absolute day (singular)", "2024-06-15", "day=1", "2024-06-01", "Pass", "Ahmad"],
    ["TC_RD11", "relativedelta", "ECP", "Large offset 100 years", "2024-01-01", "years=+100", "2124-01-01", "Pass", "Ahmad"],
    ["TC_RD12", "relativedelta", "ECP", "Delta between two dates", "dt1=2024-01-10, dt2=2025-03-15", "delta(dt2,dt1)", "years=1,months=2,days=5", "Pass", "Ahmad"],
    ["TC_RD13", "relativedelta", "BVA", "Year boundary crossing", "2024-12-31", "days=+1", "2025-01-01", "Pass", "Ahmad"],
    ["TC_RD14", "relativedelta", "ECP", "Negation", "rd(years=1,months=2)", "-rd", "rd(years=-1,months=-2)", "Pass", "Ahmad"],
    ["TC_RD15", "relativedelta", "BVA", "Non-integer months ValueError", "N/A", "months=1.5", "ValueError", "Pass", "Ahmad"],
]

for i, row in enumerate(rd_tests, 2):
    for j, val in enumerate(row, 1):
        ws2.cell(row=i, column=j, value=val)

style_header(ws2, len(headers2))
style_rows(ws2, len(rd_tests), len(headers2))
for col in ws2.columns:
    ws2.column_dimensions[col[0].column_letter].width = 22

# ===================== Sheet 3: Rrule Tests =====================
ws3 = wb.create_sheet("Rrule Tests")
headers3 = ["Test ID", "Module", "Test Type", "Test Description", "Frequency", "Parameters", "Expected Output", "Status", "Author"]
for i, h in enumerate(headers3, 1):
    ws3.cell(row=1, column=i, value=h)

rrule_tests = [
    ["TC_RR01", "rrule", "ECP", "YEARLY frequency", "YEARLY", "count=3, dtstart=2024-01-01", "3 yearly occurrences", "Pass", "Ibrahim"],
    ["TC_RR02", "rrule", "ECP", "MONTHLY frequency", "MONTHLY", "count=4, dtstart=2024-01-15", "4 monthly occurrences", "Pass", "Ibrahim"],
    ["TC_RR03", "rrule", "ECP", "WEEKLY frequency", "WEEKLY", "count=4, dtstart=2024-01-01", "4 weekly occurrences", "Pass", "Ibrahim"],
    ["TC_RR04", "rrule", "ECP", "DAILY frequency", "DAILY", "count=5, dtstart=2024-01-01", "5 daily occurrences", "Pass", "Ibrahim"],
    ["TC_RR05", "rrule", "BVA", "count=1 single occurrence", "DAILY", "count=1", "exactly 1 occurrence", "Pass", "Ibrahim"],
    ["TC_RR06", "rrule", "BVA", "interval=2 every other day", "DAILY", "interval=2, count=4", "4 alternate-day occurrences", "Pass", "Ibrahim"],
    ["TC_RR07", "rrule", "ECP", "between() method", "DAILY", "between(jan1, jan10)", "dates between range", "Pass", "Ibrahim"],
    ["TC_RR08", "rrule", "ECP", "after() method", "DAILY", "after(jan3)", "first date after jan3", "Pass", "Ibrahim"],
    ["TC_RR09", "rrule", "ECP", "HOURLY frequency", "HOURLY", "count=4", "4 hourly occurrences", "Pass", "Ahmad"],
    ["TC_RR10", "rrule", "ECP", "byweekday single (MO)", "WEEKLY", "byweekday=MO, count=3", "3 Mondays", "Pass", "Ahmad"],
    ["TC_RR11", "rrule", "ECP", "byweekday multiple (MO,FR)", "WEEKLY", "byweekday=(MO,FR), count=5", "Mon and Fri alternating", "Pass", "Ahmad"],
    ["TC_RR12", "rrule", "BVA", "until == dtstart", "DAILY", "until=dtstart", "exactly 1 occurrence", "Pass", "Ahmad"],
    ["TC_RR13", "rrule", "BVA", "bymonthday=31 monthly", "MONTHLY", "bymonthday=31, count=7", "skips short months", "Pass", "Ahmad"],
    ["TC_RR14", "rrule", "BVA", "negative bymonthday=-1", "MONTHLY", "bymonthday=-1, count=3", "last day of each month", "Pass", "Ahmad"],
    ["TC_RR15", "rrule", "ECP", "rruleset with exdate", "DAILY", "exdate=jan3", "excludes jan3", "Pass", "Ahmad"],
]

for i, row in enumerate(rrule_tests, 2):
    for j, val in enumerate(row, 1):
        ws3.cell(row=i, column=j, value=val)

style_header(ws3, len(headers3))
style_rows(ws3, len(rrule_tests), len(headers3))
for col in ws3.columns:
    ws3.column_dimensions[col[0].column_letter].width = 22

# ===================== Sheet 4: CSV-Driven Test Summary =====================
ws4 = wb.create_sheet("CSV-Driven Tests")
headers4 = ["Test ID", "CSV File", "Module", "Description", "Data Source", "Status"]
for i, h in enumerate(headers4, 1):
    ws4.cell(row=1, column=i, value=h)

csv_tests = [
    ["TC_CSV_P", "parser_test_data.csv", "parser", "Parse dates from CSV with expected outputs", "test_data/parser_test_data.csv", "Pass"],
    ["TC_CSV_R", "relativedelta_test_data.csv", "relativedelta", "Apply relativedelta from CSV data", "test_data/relativedelta_test_data.csv", "Pass"],
    ["TC_CSV_RR", "rrule_test_data.csv", "rrule", "Generate rrule occurrences from CSV", "test_data/rrule_test_data.csv", "Pass"],
    ["TC_CSV_RP", "parser_test_data.csv", "parser", "Validate error cases from CSV", "test_data/parser_test_data.csv", "Pass"],
]

for i, row in enumerate(csv_tests, 2):
    for j, val in enumerate(row, 1):
        ws4.cell(row=i, column=j, value=val)

style_header(ws4, len(headers4))
style_rows(ws4, len(csv_tests), len(headers4))
for col in ws4.columns:
    ws4.column_dimensions[col[0].column_letter].width = 25

# Save
wb.save("c:\\Assignment-3\\docs\\test_cases.xlsx")
print("Excel file created: docs/test_cases.xlsx")
