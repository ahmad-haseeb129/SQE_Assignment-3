"""
Generate individual and combined test case Excel workbooks in assignment_docs/
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

header_font = Font(bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

def create_styled_sheet(ws, title, headers, data):
    ws.title = title
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    for row_idx, row_data in enumerate(data, 2):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(wrap_text=True, vertical='top')

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 24

# --- Ibrahim Test Cases ---
ibrahim_headers = ["Test ID", "Module", "Technique", "Description", "Input / Vector", "Expected Result", "Status"]
ibrahim_data = [
    ["TC_IB_P01", "parser", "ECP", "Parse ISO 8601 date string", "2024-01-15", "datetime(2024, 1, 15)", "Pass"],
    ["TC_IB_P02", "parser", "ECP", "Parse US slash format", "01/15/2024", "datetime(2024, 1, 15)", "Pass"],
    ["TC_IB_P03", "parser", "ECP", "Parse written month names", "January 15, 2024", "datetime(2024, 1, 15)", "Pass"],
    ["TC_IB_P04", "parser", "ECP", "Parse abbreviated month names", "Jan 15, 2024", "datetime(2024, 1, 15)", "Pass"],
    ["TC_IB_P05", "parser", "ECP", "Parse date with time string", "2024-01-15 14:30:00", "datetime(2024, 1, 15, 14, 30)", "Pass"],
    ["TC_IB_P06", "parser", "ECP", "dayfirst=True on ambiguous date", "01/02/2024, dayfirst=True", "datetime(2024, 2, 1)", "Pass"],
    ["TC_IB_P07", "parser", "BVA", "Minimum year boundary", "0001-01-01", "datetime(1, 1, 1)", "Pass"],
    ["TC_IB_P08", "parser", "BVA", "Maximum year boundary", "9999-12-31", "datetime(9999, 12, 31)", "Pass"],
    ["TC_IB_P09", "parser", "BVA", "Leap day boundary", "2024-02-29", "datetime(2024, 2, 29)", "Pass"],
    ["TC_IB_RD01", "relativedelta", "ECP", "Identity offset", "base + relativedelta()", "date(2024, 6, 15)", "Pass"],
    ["TC_IB_RD02", "relativedelta", "ECP", "Positive year offset", "date + rd(years=1)", "date(2025, 1, 15)", "Pass"],
    ["TC_IB_RD03", "relativedelta", "ECP", "Negative year offset", "date + rd(years=-1)", "date(2023, 1, 15)", "Pass"],
    ["TC_IB_RD04", "relativedelta", "ECP", "Positive month offset", "date + rd(months=3)", "date(2024, 4, 15)", "Pass"],
    ["TC_IB_RD05", "relativedelta", "ECP", "Compound offsets", "date + rd(years=1, months=2, days=5)", "date(2025, 3, 20)", "Pass"],
    ["TC_IB_RD06", "relativedelta", "ECP", "Scalar multiplication", "rd(months=2, days=3) * 3", "rd(months=6, days=9)", "Pass"],
    ["TC_IB_RD07", "relativedelta", "BVA", "Month-end leap overflow", "2024-01-31 + rd(months=1)", "2024-02-29", "Pass"],
    ["TC_IB_RD08", "relativedelta", "BVA", "Month-end non-leap overflow", "2023-01-31 + rd(months=1)", "2023-02-28", "Pass"],
    ["TC_IB_RD09", "relativedelta", "BVA", "Feb 29 leap year boundary", "2024-02-29 + rd(years=1)", "2025-02-28", "Pass"],
    ["TC_IB_CSV01", "parser", "Data-Driven", "CSV data test vectors", "assignment_tests/data/parser_cases_ibrahim.csv", "8 test vectors pass", "Pass"],
    ["TC_IB_CSV02", "relativedelta", "Data-Driven", "CSV date arithmetic vectors", "assignment_tests/data/relativedelta_cases_ibrahim.csv", "8 test vectors pass", "Pass"],
]

wb_ib = openpyxl.Workbook()
create_styled_sheet(wb_ib.active, "Ibrahim Test Cases", ibrahim_headers, ibrahim_data)
wb_ib.save(r"c:\Assignment-3\assignment_docs\test_cases_ibrahim.xlsx")

# --- Ahmad Test Cases ---
ahmad_headers = ["Test ID", "Module", "Technique", "Description", "Input / Vector", "Expected Result", "Status"]
ahmad_data = [
    ["TC_AH_P01", "parser", "ECP", "Empty string input error", "parse('')", "Raises ParserError", "Pass"],
    ["TC_AH_P02", "parser", "ECP", "Non-date conversational string", "parse('hello world')", "Raises ParserError", "Pass"],
    ["TC_AH_P03", "parser", "ECP", "None input type error", "parse(None)", "Raises TypeError", "Pass"],
    ["TC_AH_P04", "parser", "ECP", "Numeric string parsing", "parse('20240115')", "datetime(2024, 1, 15)", "Pass"],
    ["TC_AH_P05", "parser", "ECP", "Fuzzy date parsing in sentence", "sentence with date, fuzzy=True", "datetime(2024, 1, 15)", "Pass"],
    ["TC_AH_P06", "parser", "ECP", "Timezone offset parsing", "2024-01-15T14:30:00+05:00", "Preserves tzinfo (+5hr)", "Pass"],
    ["TC_AH_P07", "parser", "BVA", "Midnight boundary", "2024-01-15 00:00:00", "hour=0, min=0, sec=0", "Pass"],
    ["TC_AH_P08", "parser", "BVA", "Day-end boundary", "2024-01-15 23:59:59", "hour=23, min=59, sec=59", "Pass"],
    ["TC_AH_RR01", "rrule", "ECP", "Yearly recurrence rule", "YEARLY, count=3", "3 annual recurrences", "Pass"],
    ["TC_AH_RR02", "rrule", "ECP", "Monthly recurrence frequency", "MONTHLY, count=3", "3 monthly recurrences", "Pass"],
    ["TC_AH_RR03", "rrule", "ECP", "Weekly recurrence frequency", "WEEKLY, count=3", "3 weekly recurrences", "Pass"],
    ["TC_AH_RR04", "rrule", "ECP", "Daily recurrence frequency", "DAILY, count=4", "4 daily occurrences", "Pass"],
    ["TC_AH_RR05", "rrule", "ECP", "Hourly recurrence rule", "HOURLY, interval=2, count=3", "3 occurrences separated by 2h", "Pass"],
    ["TC_AH_RR06", "rrule", "ECP", "Single weekday filter (MO)", "byweekday=MO, count=3", "Only Mondays returned", "Pass"],
    ["TC_AH_RR07", "rrule", "ECP", "Multiple weekday filter (MO,FR)", "byweekday=(MO,FR), count=5", "Only Mon & Fri returned", "Pass"],
    ["TC_AH_RR08", "rrule", "BVA", "Recurrence count min limit (0)", "count=0", "0 occurrences", "Pass"],
    ["TC_AH_RR09", "rrule", "BVA", "Recurrence count boundary (1)", "count=1", "1 occurrence (dtstart)", "Pass"],
    ["TC_AH_RR10", "rrule", "BVA", "until equals dtstart", "until = dtstart", "1 occurrence", "Pass"],
    ["TC_AH_RR11", "rrule", "BVA", "bymonthday=31 filter", "MONTHLY, bymonthday=31", "Skips months with < 31 days", "Pass"],
    ["TC_AH_RR12", "rrule", "BVA", "bymonthday=-1 last day index", "MONTHLY, bymonthday=-1", "Dynamically yields last day of month", "Pass"],
    ["TC_AH_RR13", "rrule", "ECP", "rruleset with exdate exclusion", "rruleset + exdate", "Excludes targeted date", "Pass"],
    ["TC_AH_CSV01", "parser", "Data-Driven", "CSV error test cases", "assignment_tests/data/parser_cases_ahmad.csv", "5 error vectors pass", "Pass"],
    ["TC_AH_CSV02", "rrule", "Data-Driven", "CSV recurrence rule cases", "assignment_tests/data/rrule_cases_ahmad.csv", "8 recurrence vectors pass", "Pass"],
]

wb_ah = openpyxl.Workbook()
create_styled_sheet(wb_ah.active, "Ahmad Test Cases", ahmad_headers, ahmad_data)
wb_ah.save(r"c:\Assignment-3\assignment_docs\test_cases_ahmad.xlsx")

# Copy master test_cases.xlsx to assignment_docs/
wb_master = openpyxl.Workbook()
create_styled_sheet(wb_master.active, "Ibrahim Cases", ibrahim_headers, ibrahim_data)
create_styled_sheet(wb_master.create_sheet(), "Ahmad Cases", ahmad_headers, ahmad_data)
wb_master.save(r"c:\Assignment-3\assignment_docs\test_cases.xlsx")
wb_master.save(r"c:\Assignment-3\docs\test_cases.xlsx")

print("Generated Excel files in assignment_docs/: test_cases_ibrahim.xlsx, test_cases_ahmad.xlsx, test_cases.xlsx")
